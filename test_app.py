"""Self-check: parse a list, fill the docx, convert it, build the email. No SMTP."""
import io, os, re, zipfile
import app


def docx_text(data):
    x = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode()
    return re.sub(r"<[^>]+>", "", x)


def main():
    headers, rows = app.parse_csv(open("sample-candidates.csv", "rb").read())
    assert "email" in headers and len(rows) > 1, headers
    row = rows[0]

    if os.path.exists("admissions-list.xlsx"):        # the real list, when it is present
        xheaders, xrows = app.parse_xlsx(open("admissions-list.xlsx", "rb").read())
        assert "email" in xheaders and len(xrows) > 1, xheaders

    fixed = {"campus": "SCET", "category": "A2", "admission_type": "Open Merit",
             "status": "Provisionally Selected", "date": "01/09/2026",
             "time": "10:30 AM", "venue": "ADMISSION OFFICE SCET"}

    text = docx_text(app.fill_offer_letter(row, fixed))
    for expected in (row["Candidate Name"], row["Father Name"], row["Discipline"],
                     f"{row['Matric']} out of 1100", f"{row['Inter']} out of 520",
                     f"{row['ECAT']} out of 400.00", "A2", "01/09/2026", "10:30 AM"):
        assert expected in text, f"missing from offer letter: {expected!r}"
    assert "942 out of 1100" not in text, "sample marks left in the letter"

    from docx import Document
    doc = Document(io.BytesIO(letter := app.fill_offer_letter(row, fixed)))
    name_row = [r for r in doc.tables[1].rows if r.cells[0].text.strip() == "Candidate Name"][0]
    assert name_row.cells[1]._tc is name_row.cells[-1]._tc, "name cell not merged to the row end"
    assert name_row.cells[1].paragraphs[0].alignment is None, "merged cell kept the 00 cell's alignment"
    assert "00" not in name_row.cells[1].text
    father = [r for r in doc.tables[1].rows if r.cells[0].text.strip().startswith("Father")][0]
    assert not father.cells[-1]._tc.findall(".//" + app.W + "drawing"), "leftover drawn rule from the 00 field"
    assert not father.cells[-1]._tc.findall(".//" + app.MC + "AlternateContent")

    checklist = [t for t in doc.tables if all(len(r.cells) == 1 for r in t.rows)][0]
    assert checklist.rows[0].cells[0].text.startswith("1.  MATRIC")
    assert checklist.rows[6].cells[0].text.startswith("7.  08 PASSPORT")

    pdf = app.docx_to_pdf(letter)
    assert pdf.startswith(b"%PDF"), "LibreOffice did not return a PDF"

    msg = app.build_message(row, fixed, "Offer — $program", "Dear $name, 100% $ ok", "admin@example.com")
    assert msg["To"] == row["email"]
    assert row["Discipline"] in msg["Subject"]
    att = list(msg.iter_attachments())
    assert len(att) == 1 and att[0].get_filename() == f"Offer Letter - {row['Candidate Name']}.pdf"
    assert att[0].get_content_type() == "application/pdf"
    assert row["Candidate Name"] in msg.get_body(("plain",)).get_content()
    html = msg.get_body(("html",)).get_content()
    assert 'src="cid:logo"' in html
    logo = [p for p in msg.walk() if p.get_content_type() == "image/jpeg"]
    assert len(logo) == 1 and logo[0]["Content-ID"] == "<logo>", "logo not embedded inline"
    assert row["Candidate Name"] in html and "Sharif College" in html and fixed["time"] in html
    assert "$name" not in html and "$program" not in html

    print(f"ok — {len(rows)} candidates parsed, letter + email built for {row['Candidate Name']}")


if __name__ == "__main__":
    main()
