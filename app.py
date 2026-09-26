import csv, io, json, os, re, shutil, smtplib, subprocess, sys, tempfile, threading, time, webbrowser
from functools import lru_cache
from email.message import EmailMessage
from html import escape
from string import Template

from docx import Document
from flask import Flask, Response, jsonify, render_template, request

BUNDLED = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
# a packaged app must not write inside its own bundle, and launches with "/" as the working
# directory — so give it a predictable spot the user can actually find
HOME = (os.path.expanduser("~/Documents/Admission Offer Sender")
        if getattr(sys, "frozen", False) else BUNDLED)

app = Flask(__name__, template_folder=os.path.join(BUNDLED, "templates"),
            static_folder=os.path.join(BUNDLED, "static"))

TEMPLATE_DOCX = os.path.join(BUNDLED, "offer-letter.docx")
PREVIEW_DIR = os.path.join(HOME, "preview")
SEND_DELAY = 0.5  # ponytail: fixed pause between sends; switch to a real rate limiter if a provider complains
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
SMTP_HOSTS = {
    "gmail.com": ("smtp.gmail.com", 465),
    "outlook.com": ("smtp-mail.outlook.com", 587),
    "hotmail.com": ("smtp-mail.outlook.com", 587),
    "live.com": ("smtp-mail.outlook.com", 587),
    "yahoo.com": ("smtp.mail.yahoo.com", 465),
}
COLLEGE = "Sharif College of Engineering and Technology"
SOFFICE = shutil.which("soffice") or next(
    (p for p in ("/Applications/LibreOffice.app/Contents/MacOS/soffice",
                 r"C:\Program Files\LibreOffice\program\soffice.exe",
                 "/usr/bin/soffice") if os.path.exists(p)), "")
LOGO = os.path.join(BUNDLED, "static", "logo.jpg")
RED = "#a00601"        # the college's own red, from scet.sharif.edu.pk

ROWS = []      # list of dicts, one per candidate
HEADERS = []
BATCH = {}     # what the preview page is currently showing


# ---------- parsing ----------

def parse_csv(data):
    text = data.decode("utf-8-sig", errors="replace")
    r = csv.DictReader(io.StringIO(text))
    return [h.strip() for h in (r.fieldnames or [])], [
        {(k or "").strip(): (v or "").strip() for k, v in row.items()} for row in r
    ]


def parse_xlsx(data):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows = wb.active.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(rows)]
    out = []
    for r in rows:
        d = {h: ("" if v is None else str(v).strip()) for h, v in zip(headers, r) if h}
        if any(d.values()):
            out.append(d)
    wb.close()
    return [h for h in headers if h], out


# ---------- offer letter ----------

def set_cell(cell, value):
    """Write value into a table cell, keeping the cell's existing font."""
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = value
        for extra in p.runs[1:]:
            extra.text = ""
    else:
        p.add_run(value)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def letter_values(row, fixed):
    """Every value this candidate's letter gets filled with, keyed by the label it sits against."""
    return {
        "Candidate Name": row.get("Candidate Name", ""),
        "Father's/Guardian's Name": row.get("Father Name", ""),
        "Program": row.get("Discipline", ""),
        "Campus": fixed.get("campus", ""),
        "Admission Type": fixed.get("admission_type", ""),
        "Admission Status": fixed.get("status", ""),
        "Admission Formalities Venue": fixed.get("venue", ""),
        "Date": fixed.get("date", ""),
        "Matriculation": f"{row.get('Matric', '')} out of 1100",
        "First Year": f"{row.get('Inter', '')} out of 520",
        "Category": fixed.get("category", ""),      # sits beside Campus
        "Time": fixed.get("time", ""),              # sits beside Date
        "Test Marks": f"{row.get('ECAT', '')} out of 400.00",
    }


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"


def strip_rules(cell):
    """Remove drawn rules from a cell. Word stores each one twice — a DrawingML shape and a
    VML fallback inside mc:AlternateContent — so dropping only one leaves it rendering."""
    for tag in (MC + "AlternateContent", W + "drawing", W + "pict"):
        for el in cell._tc.findall(".//" + tag):
            el.getparent().remove(el)


def reset(cell):
    """Blank a merged cell: merge() drops empty paragraphs, so the surviving one can carry
    the wrong cell's alignment and font. Strip it back to the table default."""
    p = cell.paragraphs[0]
    p.alignment = None
    for run in list(p.runs):
        run._element.getparent().remove(run._element)
    return cell


def fill_offer_letter(row, fixed):
    """Return the offer letter for one candidate as docx bytes."""
    doc = Document(TEMPLATE_DOCX)
    values = letter_values(row, fixed)
    for table in doc.tables:
        if all(len(r.cells) == 1 for r in table.rows):   # the documents checklist — number it
            for n, r in enumerate(t for t in table.rows if t.cells[0].text.strip()):
                set_cell(r.cells[0], f"{n + 1}.  {r.cells[0].text.strip()}")
            continue
        for r in table.rows:
            label = r.cells[0].text.strip()
            if label in ("Candidate Name", "Program"):
                # these value cells are only a quarter wide — merge to the row end so long
                # program names stay on one line (and the stray "00" cell disappears with it)
                set_cell(reset(r.cells[1].merge(r.cells[-1])), values[label])
            elif label in values and values[label]:
                set_cell(r.cells[1], values[label])
            if label == "Father's/Guardian's Name":
                # a 1.2pt drawn rule, left over from the blank "00" field on the row above
                strip_rules(r.cells[-1])
            if label == "Campus":            # same row also holds Category
                set_cell(r.cells[3], values["Category"])
            elif label == "Date":            # same row also holds Time
                set_cell(r.cells[3], values["Time"])
            elif label == "Combination":     # value sits in the third cell
                set_cell(r.cells[2], values["Test Marks"])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


_convert = threading.Lock()
# our own LibreOffice profile: a copy the user has open would otherwise hold the default one,
# and every conversion would exit without writing a PDF
_profile = "file://" + tempfile.mkdtemp(prefix="offer-sender-libreoffice-")


def docx_to_pdf(data):
    """Convert the generated letter to PDF with LibreOffice, so the real document can be viewed."""
    if not os.path.exists(SOFFICE):
        raise RuntimeError("LibreOffice not found — install it to preview the letter as a document.")
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "letter.docx")
        with open(src, "wb") as f:
            f.write(data)
        with _convert:  # ponytail: one at a time — soffice shares a user profile and trips over itself
            subprocess.run([SOFFICE, f"-env:UserInstallation={_profile}",
                            "--headless", "--convert-to", "pdf", "--outdir", tmp, src],
                           check=True, capture_output=True, timeout=120)
        pdf = os.path.join(tmp, "letter.pdf")
        if not os.path.exists(pdf):
            raise RuntimeError("LibreOffice did not produce a PDF.")
        with open(pdf, "rb") as f:
            return f.read()


@lru_cache(maxsize=32)   # scrolling back over a preview should not reconvert
def cached_pdf(i, fixed_json):
    return docx_to_pdf(fill_offer_letter(ROWS[i], json.loads(fixed_json)))


def letter_name(row):
    return f"Offer Letter - {row.get('Candidate Name', 'Candidate')}.pdf"


# ---------- email ----------

def context(row, fixed):
    ctx = {k.lower().replace(" ", "_").replace("-", "_"): v for k, v in row.items()}
    ctx.update(fixed)
    ctx["name"] = row.get("Candidate Name", "")
    ctx["program"] = row.get("Discipline", "")
    ctx["college"] = COLLEGE
    return ctx


def render_html(ctx, text, logo="cid:logo"):
    """Wrap the plain-text body in a plain, well-set HTML email (tables + inline CSS for mail clients)."""
    paras = "".join(
        f'<p style="margin:0 0 18px;font-size:17px;line-height:1.7;color:#030301">{escape(p).replace(chr(10), "<br>")}</p>'
        for p in text.split("\n\n") if p.strip()
    )
    details = [("Program", ctx.get("program")), ("Campus", ctx.get("campus")),
               ("Category", ctx.get("category")), ("Admission type", ctx.get("admission_type")),
               ("Status", ctx.get("status")), ("Reporting date", ctx.get("date")),
               ("Reporting time", ctx.get("time")), ("Venue", ctx.get("venue"))]
    rows = "".join(
        f'<tr>'
        f'<td style="padding:11px 0;border-bottom:1px solid #e3e5e9;font-size:15px;'
        f'color:#5f677a;white-space:nowrap;width:42%;vertical-align:top">{escape(k)}</td>'
        f'<td style="padding:11px 0;border-bottom:1px solid #e3e5e9;font-size:16px;font-weight:600;color:#030301">{escape(v)}</td>'
        f'</tr>'
        for k, v in details if v
    )
    return f"""<!doctype html>
<html><body style="margin:0;padding:0;background:#f7f8f9">
<table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="background:#f7f8f9;padding:28px 12px">
<tr><td align="center">
  <table role="presentation" width="600" cellpadding="0" cellspacing="0"
         style="width:620px;max-width:100%;background:#ffffff;border:1px solid #e3e5e9;border-radius:12px;
                font-family:Helvetica,Arial,sans-serif">

    <tr><td style="padding:26px 34px 22px">
      <table role="presentation" cellpadding="0" cellspacing="0"><tr>
        <td style="padding-right:18px;vertical-align:middle">
          <img src="{logo}" alt="" width="62" height="86"
               style="display:block;width:62px;height:86px;border:0">
        </td>
        <td style="vertical-align:middle">
          <div style="font-size:21px;font-weight:bold;color:#030301;line-height:1.3">{escape(COLLEGE)}</div>
          <div style="font-size:15px;color:#5f677a;margin-top:7px">
            Admission Office &middot; Affiliated with UET Lahore</div>
        </td>
      </tr></table>
    </td></tr>
    <tr><td style="height:3px;background:{RED};font-size:0;line-height:0">&nbsp;</td></tr>

    <tr><td style="padding:30px 34px 4px">
      <div style="font-size:21px;font-weight:bold;color:{RED};margin-bottom:22px">
        Provisional Admission Offer</div>
      {paras}
    </td></tr>

    <tr><td style="padding:10px 34px 34px">
      <table role="presentation" width="100%" cellpadding="0" cellspacing="0"
             style="border-top:2px solid #030301">{rows}</table>
    </td></tr>

    <tr><td style="border-top:1px solid #e3e5e9;padding:20px 34px;font-size:14px;line-height:1.6;color:#5f677a">
      Admission Office &middot; {escape(COLLEGE)}<br>
      This is an automated message. Please do not reply to this email.
    </td></tr>

  </table>
</td></tr></table>
</body></html>"""


def render(row, fixed, subject, body):
    """Return (subject, plain text, html) for one candidate."""
    ctx = context(row, fixed)
    text = Template(body).safe_substitute(ctx)
    return Template(subject).safe_substitute(ctx), text, render_html(ctx, text)


def build_message(row, fixed, subject, body, sender):
    subject, text, html = render(row, fixed, subject, body)
    msg = EmailMessage()
    msg["From"] = f"{COLLEGE} Admission Office <{sender}>"
    msg["To"] = row["email"]
    msg["Subject"] = subject
    msg.set_content(text)
    msg.add_alternative(html, subtype="html")
    with open(LOGO, "rb") as f:      # inline, so mail clients cannot block it
        msg.get_payload()[1].add_related(f.read(), maintype="image", subtype="jpeg", cid="<logo>")
    msg.add_attachment(docx_to_pdf(fill_offer_letter(row, fixed)),
                       maintype="application", subtype="pdf", filename=letter_name(row))
    return msg


APP_PASSWORD = re.compile(r"^[a-z]{4}( [a-z]{4}){3}$")


def connect(sender, password, host, port):
    # Google shows app passwords as "abcd efgh ijkl mnop"; the spaces are display only
    if APP_PASSWORD.match(password):
        password = password.replace(" ", "")
    domain = sender.rsplit("@", 1)[-1].lower()
    host = host or SMTP_HOSTS.get(domain, ("", 0))[0]
    port = int(port or SMTP_HOSTS.get(domain, ("", 465))[1])
    if not host:
        raise ValueError(
            "Unknown mail provider for this address — fill in the SMTP host and port. "
            "Google Workspace (your own domain): smtp.gmail.com, port 465. "
            "Microsoft 365: smtp.office365.com, port 587.")
    try:
        if port == 465:
            s = smtplib.SMTP_SSL(host, port, timeout=15)
        else:
            s = smtplib.SMTP(host, port, timeout=15)
            s.starttls()
    except OSError as e:      # name the address that failed — "timed out" alone says nothing
        raise ValueError(f"Could not reach {host} on port {port} — {e}. "
                         f"Check the host and port; Google Workspace uses smtp.gmail.com port 465.")
    try:
        s.login(sender, password)
    except smtplib.SMTPAuthenticationError as e:
        s.close()
        detail = " ".join(e.smtp_error.decode(errors="replace").split())
        raise ValueError(f"{host} rejected the sign-in for {sender} — {detail}")
    return s


# ---------- routes ----------

@app.get("/")
def index():
    return render_template("index.html", soffice=bool(SOFFICE))


@app.post("/upload")
def upload():
    global ROWS, HEADERS
    f = request.files.get("file")
    if not f or not f.filename:
        return jsonify(error="No file uploaded."), 400
    data = f.read()
    try:
        headers, rows = parse_xlsx(data) if f.filename.lower().endswith((".xlsx", ".xlsm")) else parse_csv(data)
    except Exception as e:
        return jsonify(error=f"Could not read {f.filename}: {e}"), 400
    if not rows:
        return jsonify(error="File has no data rows."), 400
    if "email" not in headers:
        return jsonify(error=f"No 'email' column. Found: {', '.join(headers)}"), 400
    HEADERS, ROWS = headers, rows   # only replace the loaded batch once the file is known good
    cached_pdf.cache_clear()
    return jsonify(headers=HEADERS, rows=ROWS,
                   invalid=[i for i, r in enumerate(ROWS) if not EMAIL_RE.match(r.get("email", ""))])


@app.post("/test-smtp")
def test_smtp():
    """Sign in and hang up — proves the settings work without spending a candidate on it."""
    p = request.get_json(force=True)
    try:
        connect((p.get("sender") or "").strip(), p.get("password") or "",
                p.get("host"), p.get("port")).quit()
    except Exception as e:
        return jsonify(error=str(e)), 400
    return jsonify(ok=True)


@app.post("/quit")
def quit_app():
    # a windowed app has no console to Ctrl+C and no event loop to answer Cmd+Q, so the
    # UI needs a way out; exit once the response is on its way
    threading.Timer(0.3, os._exit, [0]).start()
    return jsonify(ok=True)


@app.post("/preview/session")
def preview_session():
    """Stash what the preview page should show, then it is opened in a new tab."""
    p = request.get_json(force=True)
    BATCH.update(rows=[i for i in p.get("rows", []) if 0 <= i < len(ROWS)],
                 fixed=p.get("fixed", {}), subject=p.get("subject", ""), body=p.get("body", ""))
    return jsonify(count=len(BATCH["rows"]))


@app.get("/preview")
def preview_page():
    return render_template("preview.html")


@app.get("/preview/list")
def preview_list():
    return jsonify(items=[{"row": i, "name": ROWS[i].get("Candidate Name", ""),
                           "email": ROWS[i].get("email", ""), "program": ROWS[i].get("Discipline", "")}
                          for i in BATCH.get("rows", [])])


@app.get("/preview/<int:i>")
def preview_one(i):
    if not 0 <= i < len(ROWS):
        return jsonify(error="Candidate not found — re-upload the list."), 404
    row = ROWS[i]
    subject, text, html = render(row, BATCH.get("fixed", {}), BATCH.get("subject", ""), BATCH.get("body", ""))
    html = html.replace("cid:logo", "/static/logo.jpg")   # the browser cannot resolve a cid: reference
    return jsonify(to=row.get("email", ""), subject=subject, text=text, html=html,
                   filename=letter_name(row))


@app.get("/preview/<int:i>/pdf")
def preview_pdf(i):
    if not 0 <= i < len(ROWS):
        return jsonify(error="Candidate not found — re-upload the list."), 404
    try:
        pdf = cached_pdf(i, json.dumps(BATCH.get("fixed", {}), sort_keys=True))
    except Exception as e:
        return jsonify(error=str(e)), 501
    return Response(pdf, mimetype="application/pdf",
                    headers={"Content-Disposition": f'inline; filename="{letter_name(ROWS[i])}"'})


@app.post("/send")
def send():
    p = request.get_json(force=True)
    sender, password = (p.get("sender") or "").strip(), p.get("password") or ""
    indexes = [i for i in p.get("rows", []) if 0 <= i < len(ROWS)]
    if not EMAIL_RE.match(sender) or not password:
        return jsonify(error="Sender email and password are required."), 400
    if not indexes:
        return jsonify(error="No candidates selected."), 400
    if not os.path.exists(SOFFICE):
        return jsonify(error="LibreOffice is required to build the PDF letters — install it and restart."), 400

    fixed, subject, body = p.get("fixed", {}), p.get("subject", ""), p.get("body", "")
    dry_run = bool(p.get("dry_run"))
    if dry_run:                                  # fresh folder each dry run
        os.makedirs(PREVIEW_DIR, exist_ok=True)
        for f in os.listdir(PREVIEW_DIR):
            if f.endswith(".eml"):
                os.remove(os.path.join(PREVIEW_DIR, f))

    def stream():
        try:
            smtp = None if dry_run else connect(sender, password, p.get("host"), p.get("port"))
        except Exception as e:
            yield json.dumps({"fatal": str(e)}) + "\n"
            return
        try:
            for n, i in enumerate(indexes):
                row = ROWS[i]
                try:
                    if not EMAIL_RE.match(row.get("email", "")):
                        raise ValueError("invalid email address")
                    msg = build_message(row, fixed, subject, body, sender)
                    if smtp:
                        smtp.send_message(msg)
                    else:                        # dry run: drop the whole email on disk to inspect
                        safe = re.sub(r"[^\w .-]", "_", row.get("Candidate Name", "candidate"))
                        with open(os.path.join(PREVIEW_DIR, f"{i + 1:03d} {safe}.eml"), "wb") as fh:
                            fh.write(msg.as_bytes())
                    result = {"row": i, "ok": True}
                except Exception as e:
                    result = {"row": i, "ok": False, "error": str(e)}
                yield json.dumps(result) + "\n"
                if smtp and n < len(indexes) - 1:
                    time.sleep(SEND_DELAY)
            if dry_run:
                yield json.dumps({"folder": PREVIEW_DIR}) + "\n"
        finally:
            if smtp:
                smtp.quit()

    return Response(stream(), mimetype="application/x-ndjson")


def main():
    port = int(os.environ.get("PORT", 8765))
    url = f"http://127.0.0.1:{port}"
    if not os.environ.get("NO_BROWSER"):
        threading.Timer(1.0, webbrowser.open, [url]).start()
    print(f"Admission Offer Sender running — {url}   (close this window to quit)")
    app.run(host="127.0.0.1", port=port, threaded=True)


if __name__ == "__main__":
    main()
